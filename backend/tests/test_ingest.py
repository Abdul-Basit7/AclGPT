"""Batched, rate-limit-tolerant indexing — the thing that makes large files work."""

import io
import json

import pytest

from app.config import settings
from app.services import extract, ingest


@pytest.fixture(autouse=True)
def fast_limiter(monkeypatch):
    """Remove the pacing delay so tests are not wall-clock bound.

    These tests cover the *remote* provider path -- batching, pacing and
    retries all exist to survive a quota. The local path is covered separately
    below.
    """
    monkeypatch.setattr(settings, "embedding_provider", "google")
    monkeypatch.setattr(ingest._limiter, "_interval_per_unit", 0.0)
    monkeypatch.setattr(ingest.time, "sleep", lambda _seconds: None)


def make_chunks(count):
    from langchain_core.documents import Document as LCDocument

    chunks = [
        LCDocument(page_content=f"chunk {i}", metadata={"document_id": 1, "page": 1})
        for i in range(count)
    ]
    return chunks, [f"1-{i}" for i in range(count)]


def test_embedding_is_split_into_batches(monkeypatch):
    monkeypatch.setattr(settings, "embed_batch_size", 64)
    calls = []
    monkeypatch.setattr(
        ingest.vectorstore,
        "add_documents",
        lambda cid, docs, ids: calls.append(len(docs)),
    )

    chunks, ids = make_chunks(500)
    progress = []
    ingest._embed_in_batches(1, chunks, ids, progress.append)

    assert calls == [64] * 7 + [52], f"expected paced batches, got {calls}"
    assert sum(calls) == 500
    assert progress[-1] == 500
    assert progress == sorted(progress), "progress must only move forward"


def test_rate_limit_is_retried_then_succeeds(monkeypatch):
    monkeypatch.setattr(settings, "embed_batch_size", 100)
    attempts = {"n": 0}

    def flaky(collection_id, docs, ids):
        attempts["n"] += 1
        if attempts["n"] < 3:
            raise RuntimeError(
                "429 You exceeded your current quota. Please retry in 8.99s"
            )

    monkeypatch.setattr(ingest.vectorstore, "add_documents", flaky)

    chunks, ids = make_chunks(100)
    done = []
    ingest._embed_in_batches(1, chunks, ids, done.append)

    assert attempts["n"] == 3, "should retry the rate-limited batch until it lands"
    assert done == [100]


def test_rate_limit_gives_up_after_max_retries(monkeypatch):
    monkeypatch.setattr(settings, "embed_batch_size", 100)
    monkeypatch.setattr(settings, "embed_max_retries", 2)

    def always_limited(collection_id, docs, ids):
        raise RuntimeError("429 quota exceeded")

    monkeypatch.setattr(ingest.vectorstore, "add_documents", always_limited)

    chunks, ids = make_chunks(10)
    with pytest.raises(RuntimeError):
        ingest._embed_in_batches(1, chunks, ids, lambda _n: None)


def test_non_rate_limit_errors_are_not_retried(monkeypatch):
    monkeypatch.setattr(settings, "embed_batch_size", 100)
    attempts = {"n": 0}

    def boom(collection_id, docs, ids):
        attempts["n"] += 1
        raise ValueError("malformed request")

    monkeypatch.setattr(ingest.vectorstore, "add_documents", boom)

    chunks, ids = make_chunks(10)
    with pytest.raises(ValueError):
        ingest._embed_in_batches(1, chunks, ids, lambda _n: None)
    assert attempts["n"] == 1, "a non-retryable error must fail fast"


def test_limiter_charges_per_document_not_per_call():
    """The provider quota counts documents, so a batch of 64 must cost 64 slots."""
    limiter = ingest._RateLimiter(per_minute=60)  # 1 second per document
    limiter.reserve(10)
    first = limiter._next_allowed
    limiter.reserve(10)
    assert limiter._next_allowed - first == pytest.approx(10.0), (
        "a 10-document batch must reserve 10 seconds at 60/min, not 1"
    )


def test_retry_delay_honours_the_providers_suggestion():
    exc = RuntimeError("429 quota exceeded. Please retry in 8.99222092s")
    assert ingest._retry_delay(exc, 0) == pytest.approx(9.99222092)
    # Without a hint, back off exponentially but stay bounded.
    assert ingest._retry_delay(RuntimeError("429"), 3) == 8.0
    assert ingest._retry_delay(RuntimeError("429"), 20) == 120.0


def test_partial_embedding_is_rolled_back_on_failure(monkeypatch, tmp_path):
    """A half-indexed document must not leave orphaned vectors behind."""
    monkeypatch.setattr(settings, "embed_batch_size", 10)
    monkeypatch.setattr(settings, "embed_max_retries", 0)

    added, deleted = [], []

    def add(collection_id, docs, ids):
        if len(added) >= 2:
            raise RuntimeError("429 quota exceeded")
        added.extend(ids)

    monkeypatch.setattr(ingest.vectorstore, "add_documents", add)
    monkeypatch.setattr(
        ingest.vectorstore, "delete_vectors", lambda cid, ids: deleted.extend(ids)
    )

    # Drive the real ingest_document path so the rollback branch is exercised.
    from app.database import SessionLocal
    from app.models import Collection, Document, User

    db = SessionLocal()
    user = User(email=f"rollback{len(added)}@example.com", password_hash=None)
    db.add(user); db.commit(); db.refresh(user)
    collection = Collection(owner_id=user.id, name="c")
    db.add(collection); db.commit(); db.refresh(collection)
    document = Document(
        collection_id=collection.id, filename="big.txt", status="pending", size_bytes=1
    )
    db.add(document); db.commit(); db.refresh(document)
    doc_id = document.id
    db.close()

    ingest.stored_path(doc_id, "big.txt").write_text("word " * 4000, encoding="utf-8")
    ingest.ingest_document(doc_id)

    db = SessionLocal()
    try:
        stored = db.get(Document, doc_id)
        assert stored.status == "failed"
        assert stored.vector_ids == [], "failed docs must not claim vectors"
        assert stored.chunks_embedded == 0
    finally:
        db.close()

    assert deleted, "partially embedded chunks must be deleted from the index"
    assert set(deleted) <= set(added)


def test_rate_limit_failures_get_an_actionable_message():
    message = ingest._friendly_error(RuntimeError("429 quota exceeded for embed"))
    assert "rate limit" in message.lower()
    assert "free-tier" in message or "paid key" in message


def test_upload_limit_is_generous_enough_for_real_documents():
    assert settings.max_upload_mb >= 100, "large PDFs were being rejected outright"


# --- new file formats ---


def write(tmp_path, name, content):
    path = tmp_path / name
    if isinstance(content, bytes):
        path.write_bytes(content)
    else:
        path.write_text(content, encoding="utf-8")
    return path


def test_csv_becomes_labelled_records(tmp_path):
    path = write(
        tmp_path,
        "data.csv",
        "name,role,city\nAda,engineer,London\nGrace,admiral,Arlington\n",
    )
    (page,) = extract.extract_pages(path, "data.csv")
    text = page[1]
    assert "Columns: name, role, city" in text
    # Each row keeps its column names, so a retrieved chunk is self-describing.
    assert "name: Ada | role: engineer | city: London" in text
    assert "name: Grace | role: admiral | city: Arlington" in text


def test_tsv_and_semicolon_delimiters_are_detected(tmp_path):
    tsv = write(tmp_path, "d.tsv", "a\tb\n1\t2\n")
    assert "a: 1 | b: 2" in extract.extract_pages(tsv, "d.tsv")[0][1]

    semi = write(tmp_path, "d.csv", "a;b\n1;2\n")
    assert "a: 1 | b: 2" in extract.extract_pages(semi, "d.csv")[0][1]


def test_json_structure_is_preserved(tmp_path):
    payload = {"patient": {"age": 31, "graft": "hamstring"}, "cleared": False}
    path = write(tmp_path, "case.json", json.dumps(payload))
    text = extract.extract_pages(path, "case.json")[0][1]
    assert '"graft": "hamstring"' in text
    assert '"age": 31' in text


def test_malformed_json_still_ingests_as_text(tmp_path):
    path = write(tmp_path, "broken.json", "{not valid json,,,}")
    text = extract.extract_pages(path, "broken.json")[0][1]
    assert "not valid json" in text


def test_jsonl_records_are_separated(tmp_path):
    path = write(
        tmp_path,
        "log.jsonl",
        '{"event":"upload","ok":true}\n{"event":"query","ok":false}\n',
    )
    text = extract.extract_pages(path, "log.jsonl")[0][1]
    assert "Record 1:" in text and "Record 2:" in text
    assert '"event": "upload"' in text


def test_yaml_is_normalised(tmp_path):
    path = write(tmp_path, "conf.yaml", "server:\n  port: 8000\n  debug: true\n")
    text = extract.extract_pages(path, "conf.yaml")[0][1]
    assert "port: 8000" in text


def test_html_tags_and_scripts_are_stripped(tmp_path):
    path = write(
        tmp_path,
        "page.html",
        "<html><head><style>body{color:red}</style>"
        "<script>alert('x')</script></head>"
        "<body><h1>ACL recovery</h1><p>Takes months.</p></body></html>",
    )
    text = extract.extract_pages(path, "page.html")[0][1]
    assert "ACL recovery" in text
    assert "Takes months." in text
    assert "alert" not in text, "script content must not be indexed"
    assert "color:red" not in text, "style content must not be indexed"
    assert "<h1>" not in text


def test_xlsx_sheets_are_labelled(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Timeline"
    sheet.append(["phase", "weeks"])
    sheet.append(["strength", 12])
    buffer = io.BytesIO()
    workbook.save(buffer)

    path = write(tmp_path, "plan.xlsx", buffer.getvalue())
    text = extract.extract_pages(path, "plan.xlsx")[0][1]
    assert "Sheet: Timeline" in text
    assert "phase: strength | weeks: 12" in text


def test_large_tables_are_truncated_not_dropped(tmp_path, monkeypatch):
    monkeypatch.setattr(extract, "MAX_TABLE_ROWS", 5)
    rows = "\n".join(f"{i},value{i}" for i in range(50))
    path = write(tmp_path, "big.csv", f"id,value\n{rows}\n")
    text = extract.extract_pages(path, "big.csv")[0][1]
    assert "truncated after 5 rows" in text
    assert "id: 0" in text


@pytest.mark.parametrize(
    "name",
    ["a.csv", "a.tsv", "a.json", "a.jsonl", "a.yaml", "a.yml", "a.html", "a.xml", "a.xlsx"],
)
def test_new_formats_are_advertised_as_supported(name):
    assert extract.is_supported(name)


def test_unsupported_types_are_still_refused():
    assert not extract.is_supported("malware.exe")
    assert not extract.is_supported("archive.zip")


def test_daily_quota_is_not_retried(monkeypatch):
    """A per-day cap cannot clear by waiting, so retrying only wastes minutes."""
    monkeypatch.setattr(settings, "embed_batch_size", 8)
    attempts = {"n": 0}

    daily = (
        '429 You exceeded your current quota. quota_id: '
        '"EmbedContentRequestsPerDayPerUserPerProjectPerModel-FreeTier" '
        "Please retry in 30s"
    )

    def exhausted(collection_id, docs, ids):
        attempts["n"] += 1
        raise RuntimeError(daily)

    monkeypatch.setattr(ingest.vectorstore, "add_documents", exhausted)
    chunks, ids = make_chunks(8)
    with pytest.raises(RuntimeError):
        ingest._embed_in_batches(1, chunks, ids, lambda _n: None)

    assert attempts["n"] == 1, "a daily cap must fail immediately, not retry"


def test_daily_and_minute_quotas_are_distinguished():
    daily = RuntimeError('quota_id: "EmbedContentRequestsPerDayPerUser-FreeTier"')
    minute = RuntimeError("429 quota exceeded. Please retry in 8.9s")

    assert ingest._is_daily_quota(daily)
    assert not ingest._is_rate_limit(daily), "daily caps are not retryable"
    assert "daily quota" in ingest._friendly_error(daily)
    assert "1,000" in ingest._friendly_error(daily)

    assert not ingest._is_daily_quota(minute)
    assert ingest._is_rate_limit(minute), "per-minute limits are retryable"
    assert "per-minute" in ingest._friendly_error(minute)


# --- local embeddings ---


def test_local_provider_is_not_paced(monkeypatch):
    """A local model has no quota, so reserving slots would only add delay."""
    monkeypatch.setattr(settings, "embedding_provider", "local")
    monkeypatch.setattr(settings, "local_embed_batch_size", 128)
    monkeypatch.setattr(ingest.vectorstore, "add_documents", lambda *a: None)

    reserved = []
    monkeypatch.setattr(
        ingest._limiter, "reserve", lambda units=1: reserved.append(units)
    )

    chunks, ids = make_chunks(300)
    ingest._embed_in_batches(1, chunks, ids, lambda _n: None)
    assert reserved == [], "local embedding must not go through the rate limiter"


def test_local_provider_uses_large_batches(monkeypatch):
    monkeypatch.setattr(settings, "embedding_provider", "local")
    monkeypatch.setattr(settings, "local_embed_batch_size", 128)
    calls = []
    monkeypatch.setattr(
        ingest.vectorstore,
        "add_documents",
        lambda cid, docs, ids: calls.append(len(docs)),
    )

    chunks, ids = make_chunks(300)
    ingest._embed_in_batches(1, chunks, ids, lambda _n: None)
    assert calls == [128, 128, 44]


def test_batch_size_follows_the_provider(monkeypatch):
    monkeypatch.setattr(settings, "embed_batch_size", 8)
    monkeypatch.setattr(settings, "local_embed_batch_size", 256)

    monkeypatch.setattr(settings, "embedding_provider", "local")
    assert settings.embeddings_are_local
    assert settings.effective_embed_batch_size == 256

    monkeypatch.setattr(settings, "embedding_provider", "google")
    assert not settings.embeddings_are_local
    assert settings.effective_embed_batch_size == 8
